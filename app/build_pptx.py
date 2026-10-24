import io

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches

SLIDE_WIDTH = Inches(13.333)
BACKGROUND_COLOR = RGBColor(0x00, 0x00, 0x00)
NOTES_SEPARATOR = "\n\n———\n"


def build_presentation(slides, output_path):
    first_width, first_height = _image_size(slides[0]["png"])

    prs = Presentation()
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = int(SLIDE_WIDTH * first_height / first_width)
    layout = prs.slide_layouts[6]

    for content in slides:
        slide = prs.slides.add_slide(layout)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = BACKGROUND_COLOR
        _add_picture_centered(slide, content["png"], prs.slide_width, prs.slide_height)

        notes = NOTES_SEPARATOR.join(part for part in (content["notes"], content["text"]) if part)
        if notes:
            slide.notes_slide.notes_text_frame.text = notes

    prs.save(output_path)


def _image_size(image_bytes):
    return Image.open(io.BytesIO(image_bytes)).size


def _add_picture_centered(slide, image_bytes, area_width, area_height):
    img_width_px, img_height_px = _image_size(image_bytes)
    aspect = img_width_px / img_height_px

    width = area_width
    height = int(width / aspect)
    if height > area_height:
        height = area_height
        width = int(height * aspect)

    left = (area_width - width) // 2
    top = (area_height - height) // 2

    slide.shapes.add_picture(io.BytesIO(image_bytes), left, top, width=width, height=height)
